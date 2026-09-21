#!/usr/bin/env python3
"""Local disk indexer. NDJSON requests and responses over stdio; no network access."""
import argparse
import json
import os
import re
import sqlite3
import sys
import threading
import time
from pathlib import Path

TEXT_EXT = {'.txt', '.md', '.markdown', '.rst', '.log', '.csv', '.tsv', '.json', '.jsonl', '.xml', '.yaml', '.yml', '.toml', '.ini', '.cfg', '.conf', '.html', '.htm', '.css', '.js', '.jsx', '.ts', '.tsx', '.py', '.rb', '.go', '.rs', '.java', '.kt', '.swift', '.c', '.cc', '.cpp', '.h', '.hpp', '.sh', '.bash', '.zsh', '.sql', '.tex', '.org', '.ipynb'}
DOC_EXT = TEXT_EXT | {'.pdf', '.docx'}
MAX_FILE_BYTES = 40 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
SKIP_UNIX = {'/proc', '/sys', '/dev', '/run', '/snap'}


def extract(path):
    ext = path.suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(str(path), strict=False)
        parts, length = [], 0
        for page in reader.pages:
            chunk = page.extract_text() or ''
            parts.append(chunk)
            length += len(chunk)
            if length >= MAX_TEXT_CHARS:
                break
        return '\n'.join(parts)[:MAX_TEXT_CHARS]
    if ext == '.docx':
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(c.text for c in row.cells)
        return '\n'.join(parts)[:MAX_TEXT_CHARS]
    with path.open('rb') as f:
        raw = f.read(min(MAX_FILE_BYTES, MAX_TEXT_CHARS * 4))
    if b'\0' in raw[:4096]:
        return ''
    return raw.decode('utf-8-sig', errors='replace')[:MAX_TEXT_CHARS]


class Index:
    def __init__(self, db_path, emit):
        Path(db_path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(db_path, check_same_thread=False)
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.execute('PRAGMA busy_timeout=5000')
        self.db.execute('CREATE TABLE IF NOT EXISTS files (id INTEGER PRIMARY KEY, path TEXT UNIQUE, mtime_ns INTEGER, size INTEGER, kind TEXT, modified REAL, preview TEXT, scan_id TEXT)')
        self.db.execute('CREATE VIRTUAL TABLE IF NOT EXISTS texts USING fts5(path UNINDEXED, title, body, tokenize="unicode61")')
        self.db.execute('CREATE INDEX IF NOT EXISTS idx_scan ON files(scan_id)')
        self.db.commit()
        self.lock = threading.RLock()
        self.emit = emit
        self.cancel = threading.Event()
        self.thread = None

    def search(self, query, limit=30):
        terms = re.findall(r'[\w]{2,}', query, re.UNICODE)[:12]
        if not terms:
            return []
        expr = ' OR '.join('"' + t.replace('"', '""') + '"' for t in terms)
        with self.lock:
            rows = self.db.execute('SELECT f.path, f.kind, f.modified, f.preview, snippet(texts, 2, "", "", " … ", 35), bm25(texts) FROM texts JOIN files f ON f.id=texts.rowid WHERE texts MATCH ? ORDER BY bm25(texts) LIMIT ?', (expr, min(max(int(limit), 1), 100))).fetchall()
        return [{'path': p, 'kind': k, 'modified': m, 'preview': snip or preview, 'score': -score} for p, k, m, preview, snip, score in rows]

    def roots(self):
        if os.name == 'nt':
            import string
            return [f'{c}:\\' for c in string.ascii_uppercase if os.path.exists(f'{c}:\\')]
        if sys.platform == 'darwin':
            return ['/', *[str(p) for p in Path('/Volumes').iterdir() if p.is_dir()]]
        return ['/']

    def start(self, roots):
        if self.thread and self.thread.is_alive():
            return False
        roots = [str(Path(p).resolve()) for p in roots if os.path.isdir(p)]
        if not roots:
            return False
        self.cancel.clear()
        self.thread = threading.Thread(target=self._scan, args=(roots,), daemon=True)
        self.thread.start()
        return True

    def _scan(self, roots):
        scanned = indexed = unchanged = errors = 0
        started = time.time()
        token = str(time.time_ns())
        for root in roots:
            if self.cancel.is_set():
                break
            for base, dirs, names in os.walk(root, topdown=True, followlinks=False):
                if self.cancel.is_set():
                    break
                dirs[:] = [d for d in dirs if not os.path.islink(os.path.join(base, d)) and os.path.join(base, d) not in SKIP_UNIX]
                for name in names:
                    if self.cancel.is_set():
                        break
                    path = Path(base) / name
                    if path.suffix.lower() not in DOC_EXT or path.is_symlink():
                        continue
                    scanned += 1
                    try:
                        stat = path.stat()
                        if not stat.st_size or stat.st_size > MAX_FILE_BYTES:
                            continue
                        with self.lock:
                            old = self.db.execute('SELECT id,mtime_ns,size FROM files WHERE path=?', (str(path),)).fetchone()
                        if old and old[1] == stat.st_mtime_ns and old[2] == stat.st_size:
                            with self.lock:
                                self.db.execute('UPDATE files SET scan_id=? WHERE id=?', (token, old[0]))
                            unchanged += 1
                        else:
                            content = extract(path)
                            if not content.strip():
                                continue
                            with self.lock:
                                if old:
                                    self.db.execute('DELETE FROM texts WHERE rowid=?', (old[0],))
                                    self.db.execute('UPDATE files SET mtime_ns=?,size=?,kind=?,modified=?,preview=?,scan_id=? WHERE id=?', (stat.st_mtime_ns, stat.st_size, path.suffix.lower()[1:].upper(), stat.st_mtime, content[:350], token, old[0]))
                                    file_id = old[0]
                                else:
                                    cur = self.db.execute('INSERT INTO files(path,mtime_ns,size,kind,modified,preview,scan_id) VALUES(?,?,?,?,?,?,?)', (str(path), stat.st_mtime_ns, stat.st_size, path.suffix.lower()[1:].upper(), stat.st_mtime, content[:350], token))
                                    file_id = cur.lastrowid
                                self.db.execute('INSERT INTO texts(rowid,path,title,body) VALUES(?,?,?,?)', (file_id, str(path), name, content))
                            indexed += 1
                    except Exception:
                        errors += 1
                    if scanned % 50 == 0:
                        with self.lock:
                            self.db.commit()
                        self.emit({'event': 'progress', 'scanned': scanned, 'indexed': indexed, 'unchanged': unchanged, 'errors': errors, 'root': root})
            if not self.cancel.is_set():
                prefix = root.rstrip(os.sep) + os.sep
                with self.lock:
                    stale = self.db.execute('SELECT id FROM files WHERE (path=? OR substr(path,1,?)=?) AND scan_id != ?', (root, len(prefix), prefix, token)).fetchall()
                    for (file_id,) in stale:
                        self.db.execute('DELETE FROM texts WHERE rowid=?', (file_id,))
                        self.db.execute('DELETE FROM files WHERE id=?', (file_id,))
                    self.db.commit()
        with self.lock:
            self.db.commit()
            total = self.db.execute('SELECT count(*) FROM files').fetchone()[0]
        self.emit({'event': 'complete', 'scanned': scanned, 'indexed': indexed, 'unchanged': unchanged, 'errors': errors, 'total': total, 'cancelled': self.cancel.is_set(), 'seconds': round(time.time() - started, 1)})

    def status(self):
        with self.lock:
            total = self.db.execute('SELECT count(*) FROM files').fetchone()[0]
        return {'total': total, 'scanning': bool(self.thread and self.thread.is_alive())}


out_lock = threading.Lock()
def emit(data):
    with out_lock:
        sys.stdout.write(json.dumps(data, ensure_ascii=False) + '\n')
        sys.stdout.flush()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--db', required=True)
    args = parser.parse_args()
    idx = Index(args.db, emit)
    for line in sys.stdin:
        try:
            msg = json.loads(line)
            method = msg.get('method')
            if method == 'roots': value = idx.roots()
            elif method == 'search': value = idx.search(msg.get('query', ''), msg.get('limit', 30))
            elif method == 'scan': value = idx.start(msg.get('roots', []))
            elif method == 'cancel': idx.cancel.set(); value = True
            elif method == 'status': value = idx.status()
            else: raise ValueError('Unknown method')
            emit({'id': msg.get('id'), 'result': value})
        except Exception as e:
            emit({'id': msg.get('id'), 'error': str(e)[:300]})

if __name__ == '__main__': main()
